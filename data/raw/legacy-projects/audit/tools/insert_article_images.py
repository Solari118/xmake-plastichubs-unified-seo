from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


DOCX = Path(r"C:\SEO文章\Best_Plastics_for_CNC_Machining_WordPress_Ready_keyword_optimized.docx")
BACKUP = Path(r"C:\SEO文章\Best_Plastics_for_CNC_Machining_WordPress_Ready_keyword_optimized.before_images.docx")
IMAGE_DIR = Path(r"C:\SEO文章\Best_Plastics_for_CNC_Machining_images")

IMAGES = {
    "hero": IMAGE_DIR / "01-hero-cnc-machined-plastic-material-selection.png",
    "process": IMAGE_DIR / "03-plastic-cnc-machining-process-fixtured-stock.png",
    "nylon": IMAGE_DIR / "05-pom-vs-nylon-fit-humidity-risk.png",
    "transparent_dfm": IMAGE_DIR / "02-dfm-risk-transparent-plastic-guard.png",
    "pc_pmma": IMAGE_DIR / "04-pc-vs-pmma-transparent-part-selection.png",
    "rfq": IMAGE_DIR / "06-rfq-dfm-review-plastic-cnc-machining.png",
}

PLACEMENTS = [
    {
        "marker": "Best Plastics for CNC Machining: POM, PEEK, Nylon, PC, ABS, and PMMA",
        "image": "hero",
        "caption": "CNC-machined plastic material samples: POM, PEEK, Nylon, PC, ABS, and PMMA.",
        "alt": "CNC machined plastic material samples including POM, PEEK, Nylon, PC, ABS, and PMMA",
    },
    {
        "marker": "Plastic CNC machining materials should be reviewed before quotation, not corrected after machining.",
        "image": "process",
        "caption": "Plastic CNC machining uses controlled stock, fixturing, tooling, and chip control.",
        "alt": "Plastic CNC machining setup with plastic stock, fixture, end mill, and clean chips",
    },
    {
        "marker": "Engineering heuristic: Nylon is good for rugged functional parts.",
        "image": "nylon",
        "caption": "POM and Nylon fit review: use humidity and tolerance conditions to guide material choice.",
        "alt": "POM fit and Nylon humidity-risk comparison for CNC machined plastic spacers",
    },
    {
        "marker": "[INSERT DIAGRAM HERE:",
        "image": "transparent_dfm",
        "caption": "Transparent plastic part DFM review: fastener stress, radii, thin webs, and polishing zones.",
        "alt": "Transparent CNC machined polycarbonate guard showing fastener stress, radius, thin web, and polish zone",
        "replace_marker": True,
    },
    {
        "marker": "Engineering heuristic: Choose PMMA for clarity and appearance.",
        "image": "pc_pmma",
        "caption": "PC and PMMA selection: impact-resistant guarding versus optical clarity and cosmetic display use.",
        "alt": "Polycarbonate guard and PMMA panel comparison for transparent CNC machined plastic parts",
    },
    {
        "marker": "If the drawing only says “plastic,” “clear plastic,” or “Nylon,” the RFQ is incomplete.",
        "image": "rfq",
        "caption": "RFQ review inputs for plastic CNC machining: CAD, grade, finish, inspection, and use case.",
        "alt": "Plastic CNC machining RFQ and DFM review desk with CAD, drawings, material samples, and inspection tools",
    },
]


def paragraph_text(paragraph) -> str:
    return paragraph.text.strip()


def find_paragraph(document: Document, marker: str):
    for paragraph in document.paragraphs:
        if marker in paragraph_text(paragraph):
            return paragraph
    raise ValueError(f"Marker not found: {marker}")


def clone_paragraph_after(anchor, source):
    new_p = deepcopy(source._p)
    anchor._p.addnext(new_p)
    return anchor._parent.add_paragraph()._parent.paragraphs[-1]


def move_paragraph_after(anchor, paragraph):
    anchor._p.addnext(paragraph._p)
    return paragraph


def set_last_picture_alt(paragraph, alt: str) -> None:
    # Set title/description on the most recent inline picture for accessibility.
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("title", alt)
        doc_pr.set("descr", alt)


def add_image_after(document: Document, anchor, image_path: Path, caption: str, alt: str, width_inches: float = 6.35):
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = image_paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    set_last_picture_alt(image_paragraph, alt)
    move_paragraph_after(anchor, image_paragraph)

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True
    caption_run.font.size = Pt(9)
    move_paragraph_after(image_paragraph, caption_paragraph)
    return caption_paragraph


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def main() -> None:
    for key, image in IMAGES.items():
        if not image.exists():
            raise FileNotFoundError(f"Missing image for {key}: {image}")

    if not BACKUP.exists():
        BACKUP.write_bytes(DOCX.read_bytes())

    document = Document(DOCX)

    for placement in PLACEMENTS:
        anchor = find_paragraph(document, placement["marker"])
        if placement.get("replace_marker"):
            previous = anchor._p.getprevious()
            add_image_after(
                document,
                anchor,
                IMAGES[placement["image"]],
                placement["caption"],
                placement["alt"],
            )
            remove_paragraph(anchor)
        else:
            add_image_after(
                document,
                anchor,
                IMAGES[placement["image"]],
                placement["caption"],
                placement["alt"],
            )

    document.save(DOCX)
    print(f"Saved: {DOCX}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
