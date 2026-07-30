from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path
from typing import Any, Iterable


SKILL_DIR = Path(__file__).resolve().parents[1]
ONTOLOGY_PATH = SKILL_DIR / "references" / "service-query-ontology.json"
DEFAULT_PH_SEMRUSH = Path(
    r"C:\Users\Lostk\Documents\agent-project\data\raw\plastichubs\keywords"
    r"\PlasticHubs_SEMrush_All_Keywords_Classified_2026-07-10.xlsx"
)
DIMENSION_TOKENS = {
    "material": (
        "material", "acrylic", "pmma", "peek", "pom", "delrin", "polycarbonate",
        "nylon", "ptfe", "abs", "aluminum", "stainless", "titanium", "resin",
    ),
    "mechanism": (
        "warpage", "shrinkage", "stress", "crack", "heat", "thermal", "wear",
        "defect", "melt", "chatter", "bubble", "sink", "weld line", "creep",
    ),
    "quality": (
        "tolerance", "accuracy", "roughness", "inspection", "quality", "dimension",
        "flatness", "repeatability", "measurement",
    ),
    "finishing": (
        "surface", "finish", "polish", "coating", "paint", "anod", "plating",
        "texture", "blast", "smoothing",
    ),
    "design": (
        "dfm", "design", "wall thickness", "draft", "rib", "boss", "radius",
        "undercut", "thread", "gate",
    ),
    "business": (
        "cost", "price", "quote", "supplier", "manufacturer", "low volume",
        "small batch", "lead time", "service",
    ),
}
ALL_DIMENSIONS = ("process", "material", "mechanism", "quality", "finishing", "design", "business")


def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def number(value: Any) -> float:
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def load_ontology() -> dict[str, Any]:
    return json.loads(ONTOLOGY_PATH.read_text(encoding="utf-8"))


def canonical_headers(headers: Iterable[Any]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for raw in headers:
        if raw is None:
            continue
        key = normalize(raw)
        if key == "keyword":
            mapping["keyword"] = str(raw)
        elif key in {"volume", "max volume", "search volume"}:
            mapping.setdefault("volume", str(raw))
        elif key in {"min kd", "keyword difficulty", "kd", "keyword difficulty %"}:
            mapping.setdefault("kd", str(raw))
        elif key in {"intent", "intent set"}:
            mapping.setdefault("intent", str(raw))
        elif key == "fit decision":
            mapping["fit"] = str(raw)
        elif key == "category":
            mapping["category"] = str(raw)
        elif key == "subcategory":
            mapping["subcategory"] = str(raw)
    return mapping


def rows_from_csv(path: Path) -> Iterable[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        yield from reader


def rows_from_xlsx(path: Path) -> Iterable[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read .xlsx Semrush files") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["All Keywords"] if "All Keywords" in workbook.sheetnames else workbook.worksheets[0]
        values = sheet.iter_rows(values_only=True)
        headers = next(values)
        for row in values:
            yield dict(zip(headers, row))
    finally:
        workbook.close()


def read_keyword_rows(paths: list[Path]) -> Iterable[dict[str, Any]]:
    for path in paths:
        if not path.is_file():
            continue
        iterator = rows_from_xlsx(path) if path.suffix.lower() == ".xlsx" else rows_from_csv(path)
        for row in iterator:
            headers = canonical_headers(row.keys())
            keyword_key = headers.get("keyword")
            if not keyword_key or not row.get(keyword_key):
                continue
            yield {
                "keyword": normalize(row.get(keyword_key)),
                "volume": number(row.get(headers.get("volume", ""))),
                "kd": number(row.get(headers.get("kd", ""))),
                "intent": normalize(row.get(headers.get("intent", ""))),
                "fit": normalize(row.get(headers.get("fit", ""))),
                "category": str(row.get(headers.get("category", "")) or ""),
                "subcategory": str(row.get(headers.get("subcategory", "")) or ""),
                "source_file": str(path),
            }


def service_match(keyword: str, active: list[str], ontology: dict[str, Any]) -> tuple[str | None, int]:
    best: tuple[str | None, int] = (None, 0)
    for service_id in active:
        service = ontology["services"][service_id]
        score = 0
        for term in service["english_terms"]:
            normalized_term = normalize(term)
            if normalized_term == keyword:
                score = max(score, 12)
            elif contains_term(keyword, normalized_term):
                score = max(score, 8 + min(3, len(normalized_term.split())))
        if score > best[1]:
            best = (service_id, score)
    return best


def contains_term(text: str, term: str) -> bool:
    if re.search(r"[a-z0-9]", term):
        return bool(
            re.search(
                rf"(?<![a-z0-9]){re.escape(term)}(?![a-z0-9])",
                text,
            )
        )
    return term in text


def classify_dimension(keyword: str) -> str:
    for dimension, tokens in DIMENSION_TOKENS.items():
        if any(token in keyword for token in tokens):
            return dimension
    return "process"


def translate_keyword(keyword: str, service: dict[str, Any], ontology: dict[str, Any]) -> str:
    translations = ontology["translation_terms"]
    if keyword in translations:
        return translations[keyword]
    modifiers: list[str] = []
    matched_sources: list[str] = []
    for source, target in sorted(translations.items(), key=lambda item: len(item[0]), reverse=True):
        if any(source in previous for previous in matched_sources):
            continue
        if contains_term(keyword, source) and target not in modifiers:
            modifiers.append(target)
            matched_sources.append(source)
    base = service["wechat_terms"][0]
    if not modifiers:
        return base
    reduced: list[str] = []
    for modifier in modifiers:
        if modifier in base:
            continue
        if base in modifier:
            base = modifier
            continue
        if not any(modifier in current or current in modifier for current in reduced):
            reduced.append(modifier)
    return " ".join([base, *reduced[:3]])


def selected_services(
    brand: str, seeds: list[str], ontology: dict[str, Any], scope: str
) -> list[str]:
    brands = ["xmake", "plastichubs"] if brand == "both" else [brand]
    all_services: list[str] = []
    for brand_name in brands:
        for service_id in ontology["brands"][brand_name]["services"]:
            if service_id not in all_services:
                all_services.append(service_id)
    if scope == "portfolio" or not seeds:
        return all_services
    joined = " ".join(normalize(seed) for seed in seeds)
    matched: list[str] = []
    for service_id in all_services:
        service = ontology["services"][service_id]
        terms = service["english_terms"] + service["wechat_terms"]
        if any(normalize(term) in joined or joined in normalize(term) for term in terms):
            matched.append(service_id)
    if matched:
        return matched
    for service_id in all_services:
        service = ontology["services"][service_id]
        if any(
            normalize(query) in joined or joined in normalize(query)
            for values in service["dimensions"].values()
            for query in values
        ):
            matched.append(service_id)
    if scope == "focused" and matched:
        return matched
    return matched or all_services


def semrush_candidates(
    paths: list[Path], active: list[str], ontology: dict[str, Any], brand: str
) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    noise = tuple(ontology["noise_terms"])
    for row in read_keyword_rows(paths):
        keyword = row["keyword"]
        if row["fit"] in {"reject", "exclude", "drop"}:
            continue
        if row["intent"] == "navigational":
            continue
        if any(term in keyword for term in noise):
            continue
        if brand == "plastichubs" and any(
            term in keyword
            for term in (
                "metal 3d", "titanium 3d", "aluminum 3d", "stainless steel 3d",
                "steel 3d", "slm 3d", "metal printing",
            )
        ):
            continue
        service_id, relevance = service_match(keyword, active, ontology)
        if not service_id:
            continue
        dimension = classify_dimension(keyword)
        engineering_bonus = 4 if dimension != "process" else 0
        score = relevance + engineering_bonus + 2.2 * math.log10(row["volume"] + 1) - row["kd"] / 80
        service = ontology["services"][service_id]
        candidates.append(
            {
                "query": translate_keyword(keyword, service, ontology),
                "service": service_id,
                "service_label": service["label"],
                "dimension": dimension,
                "source": "semrush",
                "source_keyword": keyword,
                "volume": int(row["volume"]),
                "kd": row["kd"],
                "intent": row["intent"] or None,
                "score": round(score, 3),
                "evidence": row["source_file"],
            }
        )
    return candidates


def ontology_candidates(active: list[str], ontology: dict[str, Any]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for service_id in active:
        service = ontology["services"][service_id]
        for dimension, queries in service["dimensions"].items():
            for rank, query in enumerate(queries):
                result.append(
                    {
                        "query": query,
                        "service": service_id,
                        "service_label": service["label"],
                        "dimension": dimension,
                        "source": "project-ontology",
                        "source_keyword": None,
                        "volume": None,
                        "kd": None,
                        "intent": None,
                        "score": round(11 - rank * 0.15 + (1 if dimension in {"mechanism", "quality"} else 0), 3),
                        "evidence": str(ONTOLOGY_PATH),
                    }
                )
    return result


def external_candidates(
    values: list[str], active: list[str], ontology: dict[str, Any], source: str
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for value in values:
        keyword = normalize(value)
        service_id, relevance = service_match(keyword, active, ontology)
        if service_id:
            service = ontology["services"][service_id]
            query = value if re.search(r"[\u4e00-\u9fff]", value) else translate_keyword(keyword, service, ontology)
        else:
            service_id = active[0]
            service = ontology["services"][service_id]
            query = value
        result.append(
            {
                "query": query.strip(),
                "service": service_id,
                "service_label": service["label"],
                "dimension": classify_dimension(keyword),
                "source": source,
                "source_keyword": value,
                "volume": None,
                "kd": None,
                "intent": None,
                "score": 16 if source == "google-live" else 15,
                "evidence": "user input" if source == "user-seed" else "live browser research",
            }
        )
    return result


def dedupe(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    chosen: dict[str, dict[str, Any]] = {}
    for candidate in candidates:
        key = normalize(candidate["query"])
        current = chosen.get(key)
        if current is None or candidate["score"] > current["score"]:
            chosen[key] = candidate
    return list(chosen.values())


def diversify(
    candidates: list[dict[str, Any]], active: list[str], max_queries: int
) -> list[dict[str, Any]]:
    ordered = sorted(candidates, key=lambda item: (-item["score"], -(item["volume"] or 0), item["query"]))
    selected: list[dict[str, Any]] = []
    used: set[str] = set()

    def add_first(predicate) -> None:
        for item in ordered:
            key = normalize(item["query"])
            if key not in used and predicate(item):
                selected.append(item)
                used.add(key)
                return

    for service_id in active:
        before = len(selected)
        add_first(
            lambda item, service_id=service_id: (
                item["service"] == service_id
                and (item["source"] == "semrush" or item["dimension"] == "process")
            )
        )
        if len(selected) == before:
            add_first(lambda item, service_id=service_id: item["service"] == service_id)
        if len(selected) >= max_queries:
            break
    for service_id in active:
        if len(selected) >= max_queries:
            break
        add_first(
            lambda item, service_id=service_id: (
                item["service"] == service_id
                and item["source"] == "project-ontology"
                and item["dimension"] in {"mechanism", "quality", "design"}
            )
        )
    for dimension in ALL_DIMENSIONS:
        if len(selected) >= max_queries:
            break
        add_first(lambda item, dimension=dimension: item["dimension"] == dimension)
    minimum_per_service = max(2, max_queries // max(1, len(active)) - 1)
    for service_id in active:
        while (
            len(selected) < max_queries
            and sum(item["service"] == service_id for item in selected) < minimum_per_service
        ):
            before = len(selected)
            dimension_counts = {
                dimension: sum(item["dimension"] == dimension for item in selected)
                for dimension in ALL_DIMENSIONS
            }
            for dimension in sorted(ALL_DIMENSIONS, key=lambda value: dimension_counts[value]):
                add_first(
                    lambda item, service_id=service_id, dimension=dimension: (
                        item["service"] == service_id and item["dimension"] == dimension
                    )
                )
                if len(selected) > before:
                    break
            if len(selected) == before:
                break

    service_cap = max(3, math.ceil(max_queries / max(1, len(active))) + 2)
    dimension_cap = max(4, math.ceil(max_queries / 7) + 2)
    semrush_cap = math.ceil(max_queries * 0.6)
    while len(selected) < max_queries:
        added = False
        counts = {service: sum(item["service"] == service for item in selected) for service in active}
        dimension_counts = {
            dimension: sum(item["dimension"] == dimension for item in selected)
            for dimension in ALL_DIMENSIONS
        }
        semrush_count = sum(item["source"] == "semrush" for item in selected)
        for item in ordered:
            key = normalize(item["query"])
            if key in used or counts.get(item["service"], 0) >= service_cap:
                continue
            if dimension_counts.get(item["dimension"], 0) >= dimension_cap:
                continue
            if item["source"] == "semrush" and semrush_count >= semrush_cap:
                continue
            selected.append(item)
            used.add(key)
            counts[item["service"]] = counts.get(item["service"], 0) + 1
            dimension_counts[item["dimension"]] = dimension_counts.get(item["dimension"], 0) + 1
            if item["source"] == "semrush":
                semrush_count += 1
            added = True
            if len(selected) >= max_queries:
                break
        if not added:
            break

    for index, item in enumerate(selected):
        if index < max(6, max_queries // 3):
            item["tier"] = "A"
        elif item["dimension"] in {"mechanism", "quality", "material", "design"}:
            item["tier"] = "B"
        else:
            item["tier"] = "C"
    return selected


def write_outputs(
    output: Path,
    brand: str,
    scope: str,
    services: list[str],
    records: list[dict[str, Any]],
    semrush_files: list[Path],
) -> None:
    output.mkdir(parents=True, exist_ok=True)
    payload = {
        "brand": brand,
        "scope": scope,
        "services": services,
        "semrush_files": [str(path) for path in semrush_files],
        "query_count": len(records),
        "queries": records,
    }
    (output / "wechat-query-portfolio.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    lines = [
        "# Internal — Not for Publication",
        "",
        "# WeChat Query Portfolio",
        "",
        f"- Brand: {brand}",
        f"- Scope: {scope}",
        f"- Services: {', '.join(services)}",
        f"- Query count: {len(records)}",
        "- Volume/KD source: Semrush only where numeric evidence is present.",
        "- Google terms: live wording/intent evidence only; no Google volume claim.",
        "",
        "| Tier | WeChat query | Service | Dimension | Source | Volume | KD | Source keyword |",
        "|---|---|---|---|---|---:|---:|---|",
    ]
    for item in records:
        lines.append(
            "| {tier} | {query} | {service_label} | {dimension} | {source} | {volume} | {kd} | {source_keyword} |".format(
                **{
                    **item,
                    "volume": item["volume"] if item["volume"] is not None else "",
                    "kd": item["kd"] if item["kd"] is not None else "",
                    "source_keyword": item["source_keyword"] or "",
                }
            )
        )
    (output / "wechat-query-portfolio.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Expand brand, Semrush, Google, and engineering evidence into WeChat search queries."
    )
    parser.add_argument("--brand", choices=("xmake", "plastichubs", "both"), required=True)
    parser.add_argument("--scope", choices=("auto", "focused", "portfolio"), default="auto")
    parser.add_argument("--seed", action="append", default=[])
    parser.add_argument("--google-query", action="append", default=[])
    parser.add_argument("--semrush", action="append", default=[])
    parser.add_argument("--no-default-semrush", action="store_true")
    parser.add_argument("--max-queries", type=int, default=36)
    parser.add_argument("--out", required=True)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if not 6 <= args.max_queries <= 100:
        raise SystemExit("--max-queries must be between 6 and 100")
    ontology = load_ontology()
    effective_scope = args.scope
    if effective_scope == "auto":
        effective_scope = "focused" if args.seed else "portfolio"
    active = selected_services(args.brand, args.seed, ontology, effective_scope)
    semrush_files = [Path(value).resolve() for value in args.semrush]
    if (
        not args.no_default_semrush
        and args.brand in {"plastichubs", "both"}
        and DEFAULT_PH_SEMRUSH.is_file()
        and DEFAULT_PH_SEMRUSH not in semrush_files
    ):
        semrush_files.append(DEFAULT_PH_SEMRUSH)
    candidates = ontology_candidates(active, ontology)
    candidates.extend(semrush_candidates(semrush_files, active, ontology, args.brand))
    candidates.extend(external_candidates(args.seed, active, ontology, "user-seed"))
    candidates.extend(external_candidates(args.google_query, active, ontology, "google-live"))
    records = diversify(dedupe(candidates), active, args.max_queries)
    write_outputs(
        Path(args.out).resolve(),
        args.brand,
        effective_scope,
        active,
        records,
        semrush_files,
    )
    print(
        json.dumps(
            {
                "status": "ok",
                "brand": args.brand,
                "scope": effective_scope,
                "services": active,
                "queries": len(records),
                "output": str(Path(args.out).resolve()),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
